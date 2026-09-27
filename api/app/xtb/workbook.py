import io
import zipfile
from dataclasses import dataclass
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

SUMMARY_LABELS = frozenset({"Total", "Profit/loss"})
HEADER_SEARCH_ROWS = 40
MAX_ROWS = 100_000
MAX_COLS = 40
MAX_CONSECUTIVE_EMPTY_ROWS = 200
MAX_UNZIPPED_BYTES = 100 * 1024 * 1024


class XtbFormatError(Exception):
    """A file that cannot be read as an XTB export; the message is user-facing Polish."""

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass(frozen=True)
class Sheet:
    title: str
    metadata: dict[str, Any]
    pre_rows: list[tuple[Any, ...]]
    rows: list[dict[str, Any]]


def clean(value: Any) -> Any:
    if isinstance(value, str):
        return value.strip() or None
    return value


def _check_unzipped_size(content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            total = sum(info.file_size for info in archive.infolist())
    except zipfile.BadZipFile as exc:
        raise XtbFormatError("not_xlsx", "Plik nie jest poprawnym plikiem XLSX.") from exc
    if total > MAX_UNZIPPED_BYTES:
        raise XtbFormatError("file_too_large", "Plik po rozpakowaniu jest za duży.")


def _read_rows(sheet: Any) -> list[tuple[Any, ...]]:
    rows: list[tuple[Any, ...]] = []
    empty_streak = 0
    for row in sheet.iter_rows(max_col=MAX_COLS, values_only=True):
        cleaned = tuple(clean(cell) for cell in row)
        if all(cell is None for cell in cleaned):
            empty_streak += 1
            if empty_streak >= MAX_CONSECUTIVE_EMPTY_ROWS:
                break
        else:
            empty_streak = 0
        rows.append(cleaned)
        if len(rows) >= MAX_ROWS:
            break
    return rows


def open_workbook(content: bytes) -> dict[str, list[tuple[Any, ...]]]:
    _check_unzipped_size(content)
    try:
        workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        raise XtbFormatError("not_xlsx", "Plik nie jest poprawnym plikiem XLSX.") from exc
    try:
        return {sheet.title: _read_rows(sheet) for sheet in workbook.worksheets}
    finally:
        workbook.close()


def read_sheet(title: str, rows: list[tuple[Any, ...]], required_columns: frozenset[str]) -> Sheet:
    header_index = next(
        (
            index
            for index, row in enumerate(rows[:HEADER_SEARCH_ROWS])
            if required_columns <= {cell for cell in row if isinstance(cell, str)}
        ),
        None,
    )
    if header_index is None:
        raise XtbFormatError(
            "missing_columns",
            f"Zakładka „{title}” nie ma oczekiwanych kolumn: {', '.join(sorted(required_columns))}.",
        )
    header = [cell if isinstance(cell, str) else None for cell in rows[header_index]]
    pre_rows = rows[:header_index]
    metadata = {row[0]: (row[1] if len(row) > 1 else None) for row in pre_rows if row and isinstance(row[0], str)}
    data = []
    for row in rows[header_index + 1:]:
        if all(cell is None for cell in row) or row[0] in SUMMARY_LABELS:
            continue
        data.append({name: (row[i] if i < len(row) else None) for i, name in enumerate(header) if name})
    return Sheet(title, metadata, pre_rows, data)
