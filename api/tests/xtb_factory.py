"""Builds XLSX files shaped like real XTB account-history exports.

The layout (sheet names, metadata rows above each table, summary rows, the two kinds of
rows in Open Positions) was verified against real exports; values here are made up.
"""
import io
from collections.abc import Iterable
from datetime import datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

CASH_HEADER = ["Type", "Instrument", "Ticker", "Category", "Time", "Amount", "ID", "Comment", "Product", "Position ID"]
OPEN_HEADER = [
    "Product", "Instrument/Position", "Ticker", "Category", "Type", "Volume", "Value", "Current price",
    "Open price", "Open time (UTC)", "Stop Loss", "Take Profit", "Net Profit %", "Net Profit",
    "Gross Profit", "Margin", "Open Commission", "Swap", "Rollover",
]
CLOSED_HEADER = [
    "Instrument", "Ticker", "Category", "Type", "Volume", "Open Price", "Open Time (UTC)", "Close Price",
    "Close Time (UTC)", "Product", "Profit/Loss", "Gross Profit", "Purchase Value", "Sale Value",
    "Stop Loss", "Take Profit", "Commission", "Margin", "Swap", "Rollover", "Open Conversion Rate",
    "Close Conversion Rate", "Close Origin", "Position ID", "Comment",
]

PERIOD_FROM = datetime(2006, 1, 1)
PERIOD_TO = datetime(2026, 9, 26)
GENERATED = datetime(2026, 9, 26, 12, 0, 0)

Row = dict[str, Any]


def filename(prefix: str = "IKE", number: str = "56216965") -> str:
    return f"{prefix}_{number}_2006-01-01_2026-09-26.xlsx"


def cash_row(
    type_: str, amount: float, op_id: str, time: datetime, *, comment: str = "", ticker: str = "",
    instrument: str = "", category: str = "", product: str = "IKE", position_id: str = "",
) -> Row:
    return {
        "Type": type_, "Instrument": instrument, "Ticker": ticker, "Category": category, "Time": time,
        "Amount": amount, "ID": op_id, "Comment": comment, "Product": product, "Position ID": position_id,
    }


def buy_row(
    ticker: str, qty: str, price: str, amount: float, op_id: str, time: datetime, position_id: str, *,
    instrument: str = "Core S&P 500", category: str = "ETF", product: str = "IKE",
) -> Row:
    return cash_row(
        "Stock purchase", amount, op_id, time, comment=f"OPEN BUY {qty} @ {price}", ticker=ticker,
        instrument=instrument, category=category, product=product, position_id=position_id,
    )


def summary_row(
    ticker: str, name: str, volume: float, value: float, open_price: float, profit: float, *,
    category: str = "ETF", product: str = "IKE",
) -> Row:
    return {
        "Product": product, "Instrument/Position": name, "Ticker": ticker, "Category": category, "Type": "",
        "Volume": volume, "Value": value, "Current price": "", "Open price": open_price,
        "Open time (UTC)": "", "Net Profit %": 1.5, "Net Profit": profit, "Gross Profit": profit,
    }


def lot_row(
    ticker: str, position_id: str, volume: float, open_price: float, opened_at: datetime,
    current_price: float, value: float, profit: float, *, product: str = "IKE",
) -> Row:
    return {
        "Product": product, "Instrument/Position": position_id, "Ticker": ticker, "Category": "",
        "Type": "BUY", "Volume": volume, "Value": value, "Current price": current_price,
        "Open price": open_price, "Open time (UTC)": opened_at, "Net Profit %": 1.5,
        "Net Profit": profit, "Gross Profit": profit,
    }


def closed_row(
    ticker: str, position_id: str, volume: float, open_price: float, opened_at: datetime,
    close_price: float, closed_at: datetime, *, name: str = "Veolia", category: str = "STOCK", product: str = "IKE",
) -> Row:
    return {
        "Instrument": name, "Ticker": ticker, "Category": category, "Type": "BUY", "Volume": volume,
        "Open Price": open_price, "Open Time (UTC)": opened_at, "Close Price": close_price,
        "Close Time (UTC)": closed_at, "Product": product, "Commission": -1.0,
        "Open Conversion Rate": 4.3, "Close Conversion Rate": 4.25, "Close Origin": "Client", "Position ID": position_id,
    }


def build_report(
    *, account_number: str = "56216965", open_positions_account_number: str | None = None,
    product: str = "IKE", cash: Iterable[Row] = (), open_rows: Iterable[Row] = (),
    closed: Iterable[Row] = (), include_open: bool = True, include_closed: bool = True,
    include_cash: bool = True,
) -> bytes:
    cash = list(cash)
    workbook = Workbook()
    workbook.remove(workbook.active)
    if include_closed:
        sheet = workbook.create_sheet("Closed Positions")
        _period_header(sheet, account_number, "Closed Positions")
        _table(sheet, CLOSED_HEADER, closed)
        sheet.append(["Profit/loss"])
    if include_cash:
        sheet = workbook.create_sheet("Cash Operations")
        _period_header(sheet, account_number, "Cash Operations")
        _table(sheet, CASH_HEADER, cash)
        sheet.append(["Total", None, None, None, None, sum(row["Amount"] for row in cash)])
    if include_open:
        sheet = workbook.create_sheet("Open Positions")
        sheet.append(["Account number", open_positions_account_number or account_number])
        sheet.append(["Open Positions", ""])
        sheet.append(["Data as of report generated", GENERATED])
        sheet.append(["Product", "Metric", "Amount", "Currency"])
        sheet.append([product, "Open position value", 1000.0, "PLN"])
        sheet.append([product, "Open position profit", 25.0, "PLN"])
        sheet.append([])
        sheet.append(["Note", "Summary values and open positions are shown as of the report generation time"])
        _table(sheet, OPEN_HEADER, open_rows)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _period_header(sheet: Worksheet, account_number: str, title: str) -> None:
    sheet.append(["Account number", account_number])
    sheet.append([title, ""])
    sheet.append(["Date from (UTC)", PERIOD_FROM])
    sheet.append(["Date to (UTC)", PERIOD_TO])


def _table(sheet: Worksheet, header: list[str], rows: Iterable[Row]) -> None:
    sheet.append(header)
    for row in rows:
        sheet.append([row.get(column, "") for column in header])
